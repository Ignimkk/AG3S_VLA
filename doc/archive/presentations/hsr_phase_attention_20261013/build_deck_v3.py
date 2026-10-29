"""Build HSR_phase_attention_20261013_v3.pptx from slides_source_v3.md (python-pptx).

Concise deck for the advisor: cover + 4 body slides (no appendix / Q&A).
Style follows build_deck.py (colors, Noto Sans CJK KR, footer, eaLnBrk="0", table style);
build_deck.py itself is not imported or modified.

Run:  <venv with python-pptx, Pillow, lxml>/bin/python build_deck_v3.py [--rebuild]
Speaker notes are read verbatim from slides_source_v3.md ("대본" blocks).
Figures are read-only (originals are never modified).
"""
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
SRC = HERE / "slides_source_v3.md"
OUT = HERE / "HSR_phase_attention_20261013_v3.pptx"

FONT = "Noto Sans CJK KR"
INK = RGBColor(0x1E, 0x2A, 0x3A)
SUB_INK = RGBColor(0x33, 0x3D, 0x4A)
MUTED = RGBColor(0x5B, 0x65, 0x73)
ACCENT = RGBColor(0xA8, 0x4A, 0x22)      # rust, matches the C3 panel tone
TINT = RGBColor(0xFB, 0xF0, 0xEA)
HEAD = RGBColor(0x2E, 0x3B, 0x4E)
ROW_A = RGBColor(0xFF, 0xFF, 0xFF)
ROW_B = RGBColor(0xF2, 0xF4, 0xF7)
LINE = RGBColor(0xC9, 0xCE, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_BG = RGBColor(0x1E, 0x2A, 0x3A)
USE = RGBColor(0x14, 0x80, 0x55)         # "사용" (green, same family as the figures' target green)
NOUSE = RGBColor(0xC0, 0x39, 0x2B)       # "사용하지 않음"
USE_BG = RGBColor(0xE6, 0xF4, 0xEC)
NOUSE_BG = RGBColor(0xFB, 0xE9, 0xE6)
PH_PICK = RGBColor(0xE9, 0xA1, 0x00)     # phase colors used in the T59 timeline figure
PH_PLACE = RGBColor(0xE5, 0x7A, 0xA0)
PH_HOME = RGBColor(0xE0, 0x4B, 0x4B)

FIG = REPO / "AG3S/ag3s/docs/figures"
CROPS = HERE / "crops"
FIG_TIMELINE = FIG / "ppt-1013/t59-ppt-timeline-ep1809-key-N1.png"
FIG_BARS = FIG / "ppt-1013/t59-ppt-phase-bars.png"
FIG_SCENE = CROPS / "subtask-b-scene-top-row-compact.png"
FIG_C3 = CROPS / "c3_when_it_changes.png"

TOTAL = 5
_count = [0]


# ---------------------------------------------------------------- script (notes) from the source
def load_scripts(path):
    """Return [cover, s1, s2, s3, s4] speaker-note texts, verbatim from the source."""
    text = path.read_text(encoding="utf-8")
    sections = [s for s in re.split(r"\n---\n", text)]
    cover = next(s for s in sections if s.lstrip().startswith("## 표지"))
    m = re.search(r"^- 대본: (.+)$", cover, re.M)
    out = [m.group(1).strip()]
    for n in range(1, 5):
        sec = next(s for s in sections if s.lstrip().startswith(f"## {n}."))
        body = sec.split("**대본**", 1)[1].strip()
        out.append(body)
    return out


SCRIPTS = load_scripts(SRC)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers (same style as build_deck.py)
def set_font(run, size, bold=False, color=INK, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def set_bullet(par, level, size):
    pPr = par._p.get_or_add_pPr()
    indent_in = 0.22 if level == 0 else 0.2
    marl = 0.22 + 0.3 * level
    pPr.set("marL", str(int(Inches(marl))))
    pPr.set("indent", str(int(-Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buf = etree.SubElement(pPr, qn("a:buFont"))
    buf.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•" if level == 0 else "–")


def no_bullet(par):
    pPr = par._p.get_or_add_pPr()
    etree.SubElement(pPr, qn("a:buNone"))


def textbox(slide, x, y, w, h, margin=0.04, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if name:
        tb.name = name
    return tb


def add_runs(par, runs, size, color=INK, bold=False):
    """runs: str or list of (text, dict(bold=..., color=..., size=...))."""
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, opt in runs:
        r = par.add_run()
        r.text = text
        set_font(r, opt.get("size", size), bold=opt.get("bold", bold),
                 color=opt.get("color", color), italic=opt.get("italic", False))


def bullets(slide, x, y, w, h, items, size=15, space_after=8, name="Body", anchor=MSO_ANCHOR.TOP):
    """items: list of (level, runs). level -1 = plain paragraph (no bullet)."""
    tb = textbox(slide, x, y, w, h, name=name, anchor=anchor)
    tf = tb.text_frame
    first = True
    for level, runs in items:
        par = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        sz = size if level <= 0 else size - 1
        add_runs(par, runs, sz, color=INK if level <= 0 else SUB_INK)
        if level >= 0:
            set_bullet(par, level, sz)
        else:
            no_bullet(par)
        par.space_after = Pt(space_after)
        par.line_spacing = 1.12
    return tb


def image_fit(slide, path, x, y, w, h, align="center", valign="top", name=None):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + {"left": 0, "center": (w - pw) / 2, "right": w - pw}[align]
    py = y + {"top": 0, "middle": (h - ph) / 2, "bottom": h - ph}[valign]
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    if name:
        pic.name = name
    return pic, (px, py, pw, ph)


def header(slide, title, message):
    tb = textbox(slide, 0.5, 0.3, 12.33, 0.66, margin=0, anchor=MSO_ANCHOR.MIDDLE, name="Title")
    add_runs(tb.text_frame.paragraphs[0], title, 24, bold=True)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.06),
                                 Inches(12.33), Inches(0.74))
    box.name = "Message"
    box.adjustments[0] = 0.12
    box.fill.solid()
    box.fill.fore_color.rgb = TINT
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_runs(p, message, 16, color=ACCENT, bold=True)


def footer(slide, dark=False):
    _count[0] += 1
    n = _count[0]
    col = RGBColor(0xB8, 0xC0, 0xCC) if dark else RGBColor(0x8A, 0x93, 0xA0)
    tb = textbox(slide, 0.5, 7.08, 8.0, 0.3, margin=0, anchor=MSO_ANCHOR.MIDDLE, name="Footer")
    add_runs(tb.text_frame.paragraphs[0], "AG3S 진행 보고 · HSR · 2026-10-13", 9, color=col)
    tb = textbox(slide, 11.33, 7.08, 1.5, 0.3, margin=0, anchor=MSO_ANCHOR.MIDDLE, name="Page")
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_runs(p, f"{n} / {TOTAL}", 9, color=col)


def notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    lines = text.strip().split("\n")
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    for p in tf.paragraphs:
        for r in p.runs:
            set_font(r, 12)


def table(slide, x, y, w, col_w, rows, size=13, row_h=0.4, header_size=None, fills=None, name="Table",
          header_h=0.5):
    """rows: list of rows; each cell is str or list of (text, opt) runs.
    fills: optional {(r, c): RGBColor} overriding the body-row fill."""
    fills = fills or {}
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(header_h + row_h * (nr - 1)))
    shp.name = name
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(nr):
        tbl.rows[r].height = Inches(header_h if r == 0 else row_h)
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = HEAD
            else:
                cell.fill.fore_color.rgb = fills.get((r, c), ROW_A if r % 2 else ROW_B)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.08
            if r == 0:
                add_runs(p, rows[r][c], header_size or size, color=WHITE, bold=True)
            else:
                add_runs(p, rows[r][c], size)
    return shp


def new_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    return s


def arrow_head(conn):
    conn.line._get_or_add_ln().append(etree.fromstring(
        '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'type="triangle" w="med" len="med"/>'))


# ---------------------------------------------------------------- cover
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = DARK_BG
tb = textbox(s, 0.9, 1.6, 11.5, 2.6, margin=0, anchor=MSO_ANCHOR.BOTTOM, name="Title")
p = tb.text_frame.paragraphs[0]
add_runs(p, "VLA attention을 작업 단계에 맞게 쓰기", 40, color=WHITE, bold=True)
p = tb.text_frame.add_paragraph()
add_runs(p, "HSR (Hidden-state Sub-task Recognition)", 26, color=RGBColor(0xF2, 0xC6, 0xA8))
p.space_before = Pt(12)
tb = textbox(s, 0.9, 4.55, 11.5, 0.5, margin=0, name="Subtitle")
add_runs(tb.text_frame.paragraphs[0], "AG3S 진행 보고 · 2026-10-13", 18, color=RGBColor(0xC9, 0xD2, 0xDE))
cx = 0.9
for i, (lab, col) in enumerate([("pick", PH_PICK), ("place", PH_PLACE), ("home", PH_HOME)]):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(5.45), Inches(1.25), Inches(0.42))
    c.adjustments[0] = 0.5
    c.fill.solid(); c.fill.fore_color.rgb = col
    c.line.fill.background(); c.shadow.inherit = False
    c.name = f"Phase {lab}"
    tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    q = tf.paragraphs[0]; q.alignment = PP_ALIGN.CENTER
    add_runs(q, lab, 14, color=WHITE, bold=True)
    cx += 1.25
    if i < 2:
        a = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx + 0.06), Inches(5.66),
                                   Inches(cx + 0.34), Inches(5.66))
        a.line.color.rgb = RGBColor(0xC9, 0xD2, 0xDE); a.line.width = Pt(1.5)
        arrow_head(a)
        cx += 0.4
footer(s, dark=True)
notes(s, SCRIPTS[0])


# ---------------------------------------------------------------- slide 1: problem
s = new_slide()
header(s, "문제 제기 — attention만으로는 \"지금 무엇을 하는 중인지\" 모른다",
       "같은 attention이라도 작업 단계를 모르면, 들고 가는 동안과 놓은 뒤에 엉뚱한 물체를 \"지금 다루는 물체\"로 잡는다.")
_, (px, py, pw, ph) = image_fit(s, FIG_TIMELINE, 0.5, 2.02, 8.9, 4.9, align="left", name="Figure timeline")
bx = px + pw + 0.35
bullets(s, bx, 2.1, 12.83 - bx, 4.7, [
    (0, "attention만 쓰면: 들고 가는 동안 10 frame 중 8 frame에서 다른 과일을 잡거나 아무것도 못 잡음"),
    (0, "놓은 뒤에도 이미 바구니에 넣은 banana를 다시 잡으려 함"),
    (0, [("작업 단계를 함께 보면: ", {"bold": True, "color": ACCENT}),
         ("들고 가는 동안 banana 유지, 놓은 뒤 비움", {})]),
], size=16, space_after=16, anchor=MSO_ANCHOR.MIDDLE)
footer(s)
notes(s, SCRIPTS[1])


# ---------------------------------------------------------------- slide 2: HSR implementation
s = new_slide()
header(s, "HSR 구현 — VLA를 다시 학습하지 않고 내부 상태에서 작업 단계를 읽는다",
       "VLA가 동작을 만들 때 이미 계산하는 내부 값에서, 작은 분류기 하나로 pick / place / home을 읽는다.")

FLOW = [
    # (main, sub, fill, line)
    ([("카메라 영상 +", {"bold": True})], [("작업 지시문", {"bold": True})],
     RGBColor(0xEC, 0xEC, 0xEC), RGBColor(0x55, 0x55, 0x55)),
    ([("VLA", {"bold": True})], [("(π0.5, 그대로 사용)", {"size": 12})],
     RGBColor(0xD6, 0xE4, 0xF5), RGBColor(0x4A, 0x6E, 0x9A)),
    ([("내부 값", {"bold": True})], [("(layer 4 의 KV cache,", {"size": 12})], [("512 차원으로 요약)", {"size": 12})],
     RGBColor(0xFB, 0xE3, 0xD6), RGBColor(0xC0, 0x56, 0x2F)),
    ([("작은 분류기", {"bold": True})], [("(logistic regression)", {"size": 12})],
     RGBColor(0xF6, 0xD3, 0xC0), RGBColor(0xA8, 0x4A, 0x22)),
    ([("pick", {"bold": True, "color": PH_PICK}), (" / ", {}), ("place", {"bold": True, "color": PH_PLACE}),
      (" / ", {}), ("home", {"bold": True, "color": PH_HOME})],
     WHITE, RGBColor(0x55, 0x5E, 0x6B)),
]
BW, GAP, BH, BY = 2.13, 0.42, 1.12, 2.08
x = 0.5
for i, spec in enumerate(FLOW):
    *paras, fill, line = spec
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(BY), Inches(BW), Inches(BH))
    b.adjustments[0] = 0.14
    b.fill.solid(); b.fill.fore_color.rgb = fill
    b.line.color.rgb = line; b.line.width = Pt(1.5)
    b.shadow.inherit = False
    b.name = f"Flow {i + 1}"
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for j, runs in enumerate(paras):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, runs, 15)
    if i < len(FLOW) - 1:
        a = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + BW + 0.06), Inches(BY + BH / 2),
                                   Inches(x + BW + GAP - 0.06), Inches(BY + BH / 2))
        a.line.color.rgb = RGBColor(0x55, 0x5E, 0x6B); a.line.width = Pt(2)
        a.name = f"Arrow {i + 1}"
        arrow_head(a)
    x += BW + GAP

CY = BY + BH + 0.35
_, (px, py, pw, ph) = image_fit(s, FIG_SCENE, 5.45, CY, 7.38, 6.95 - CY, align="right", valign="middle",
                                name="Figure HSR scene")
bullets(s, 0.5, CY, px - 0.35 - 0.5, 6.95 - CY, [
    (0, "처음에는 VLA가 작업 단계를 문장으로 말하게 해 봤지만 실패: 우리 모델은 그런 문장을 학습한 적이 없어 "
        "들고 가는 구간 36 frame 중 30 frame에서 엉뚱한 출력"),
    (0, "대신 VLA 내부 값에서 단계를 읽는 작은 분류기를 학습"),
    (1, "VLA 재학습 없음 · 판단 시간 0.1 ms 미만"),
    (1, "학습 데이터: 시연 영상의 단계 표시 (집기 / 옮기기 / 복귀)"),
], size=15, space_after=8, anchor=MSO_ANCHOR.MIDDLE)
footer(s)
notes(s, SCRIPTS[2])


# ---------------------------------------------------------------- slide 3: HSR validation
s = new_slide()
header(s, "HSR 검증 — 단계는 잘 읽고, 그 덕분에 엉뚱한 물체를 잡는 일이 줄었다",
       "단계 판단 정확도 0.96, 그리고 작업 단계를 쓰면 들고 가는 동안 올바른 물체를 유지하는 비율이 0.59 → 0.86.")
FW = 6.75
_, (px, py, pw, ph) = image_fit(s, FIG_BARS, 12.83 - FW, 2.02, FW, 4.0, align="right", name="Figure phase bars")
TW = px - 0.3 - 0.5
B = {"bold": True}
table(s, 0.5, 2.02, TW, [TW * 0.5, TW * 0.5], [
    ["검증", "결과"],
    ["단계 판단 정확도 (실제 실행 기록)", [("0.96", B)]],
    ["들고 가는 동안 올바른 물체 유지 (16 실행, 같은 순간 비교)",
     [("attention만 0.59 → 작업 단계 사용 ", {}), ("0.86", B)]],
    ["놓은 뒤 다른 물체로 바뀐 실행 (48 실행)", [("HSR 없음 21 → HSR 사용 ", {}), ("0", B)]],
    ["작업 성공 (48 실행)", "29 대 28 — 차이 없음"],
], size=13, row_h=0.8, header_size=13, name="Validation table")
bullets(s, 0.5, 6.22, 12.33, 0.75, [
    (0, [("한계: ", {"bold": True, "color": ACCENT}),
         ("HSR은 지시문의 물체 이름이 바뀌어도, 손에서 물체를 지워도 판단이 잘 바뀌지 않음 → "
          "물체를 실제로 쥐었는지는 HSR이 아니라 gripper 측정값으로 판단", {})]),
], size=14, space_after=0, name="Limitation")
footer(s)
notes(s, SCRIPTS[3])


# ---------------------------------------------------------------- slide 4: attention use per phase
s = new_slide()
header(s, "단계별로 attention을 어떻게 쓰는가",
       "attention은 \"무엇을 잡을지\" 정할 때만 쓰고, 쥔 뒤와 놓은 뒤에는 작업 단계에 따라 물체의 역할을 바꾼다.")
_, (px, py, pw, ph) = image_fit(s, FIG_C3, 9.0, 2.02, 3.83, 4.93, align="right", name="Figure C3")
TW = px - 0.35 - 0.5
UB = {"bold": True, "color": USE}
NB = {"bold": True, "color": NOUSE}
table(s, 0.5, 2.02, TW, [1.75, (TW - 1.75) * 0.55, (TW - 1.75) * 0.45], [
    ["작업 단계", "attention 사용", "그 물체를 어떻게 다루나"],
    ["집으러 감 (pick)",
     [("사용", UB), (" — attention이 가장 강한 물체를 \"다룰 물체\"로 정함 (세 번 연속 같으면 확정)", {})],
     "손가락만 그 물체에 닿는 것 허용, 나머지는 장애물"],
    ["쥐는 순간", "attention이 흔들려도 다룰 물체를 바꾸지 않음", "그대로 유지"],
    ["들고 감 (place)", [("사용하지 않음", NB), (" — attention이 바구니 쪽으로 감", {})],
     "쥔 물체를 로봇 손의 일부로 붙여 함께 충돌 검사"],
    ["놓은 뒤 (home)", [("사용하지 않음", NB), (" — 새 물체를 정하지 않음", {})], "놓은 물체를 다시 장애물로"],
], size=13, row_h=0.9, header_size=13,
    fills={(1, 1): USE_BG, (3, 1): NOUSE_BG, (4, 1): NOUSE_BG}, name="Phase table")
bullets(s, 0.5, 6.4, TW, 0.5, [
    (0, [("누가 정하나: ", {"bold": True, "color": ACCENT}),
         ("쥔 순간은 gripper 측정값, 작업 단계는 HSR", {})]),
], size=15, space_after=0, name="Who decides")
footer(s)
notes(s, SCRIPTS[4])


# ---------------------------------------------------------------- finalize
assert _count[0] == TOTAL, _count[0]
for sl in prs.slides:
    for shp in sl.shapes:
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
# Korean: break lines between words, not inside a word (eaLnBrk="0")
for sl in prs.slides:
    for el in sl._element.iter(qn("a:p")):
        el.get_or_add_pPr().set("eaLnBrk", "0") if hasattr(el, "get_or_add_pPr") else None
if OUT.exists() and "--rebuild" not in sys.argv:
    raise SystemExit(f"{OUT} already exists; refusing to overwrite (pass --rebuild to overwrite this script's own output)")
prs.save(OUT)
print("wrote", OUT)
