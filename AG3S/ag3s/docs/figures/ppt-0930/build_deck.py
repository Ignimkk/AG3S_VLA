"""0930 세미나 deck — 사용자 템플릿(곽민기_0930 발표 자료.pptx)을 채운다.

python-pptx 로 템플릿을 열고:
  * 기존 슬라이드(표지·목차·피드백·아키텍처·모듈 5·Experiment·Future)는 그대로 두고 내용을 채운다
  * Experiment 상세 슬라이드와 'Perception 필요성' 슬라이드를 새로 만들어 제자리에 끼운다
  * 슬라이드마다 발표 대본을 speaker notes 에 넣는다 (대본 md 는 같은 SCRIPT 에서 나온다)
"""

from __future__ import annotations
from common.workspace import workspace_path as _ag3s_workspace_path

import copy
import json
import pathlib
import sys

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
ROOT = pathlib.Path(_ag3s_workspace_path(''))
FIG = ROOT / "AG3S/ag3s/docs/figures/ppt-0930"
ARCH = ROOT / "AG3S/ag3s/docs/figures/ag3s-architecture-v2/ag3s-system-architecture.png"
TEMPLATE = ROOT / "doc/곽민기_0930 발표 자료.pptx"
OUT = ROOT / "doc/곽민기_0930 발표 자료_draft.pptx"
SCRIPT_MD = ROOT / "doc/곽민기_0930 발표 대본.md"

FONT = "NanumSquare Neo Regular"
FONT_XB = "NanumSquare Neo ExtraBold"
NAVY = RGBColor(0x00, 0x20, 0x60)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x7F, 0x7E, 0x79)
RED = RGBColor(0xC0, 0x39, 0x2B)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREEN = RGBColor(0x1B, 0x8F, 0x62)
LIGHT = RGBColor(0xF1, 0xF3, 0xF8)
LINE = RGBColor(0xC8, 0xCC, 0xD6)

# 폭 추정용 (LibreOffice 가 없어 직접 잰다). NanumSquare Neo 대신 Noto CJK, 10 % 여유.
_MEASURE = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
WARN: list[str] = []


def _font(pt):
    return ImageFont.truetype(_MEASURE, int(round(pt * 4)))


def text_lines(text: str, pt: float, width_in: float) -> int:
    """width_in 폭에 pt 크기로 쓸 때의 줄 수 (문자 단위 줄바꿈, 10 % 여유)."""
    f = _font(pt)
    limit = width_in * 72 * 4 / 1.10
    n = 0
    for para in text.split("\n"):
        w, lines = 0.0, 1
        for ch in para:
            cw = f.getlength(ch)
            if w + cw > limit:
                lines, w = lines + 1, cw
            else:
                w += cw
        n += lines
    return n


def line_h(pt, spacing=1.2):
    return pt * spacing / 72.0


# ── low-level text ─────────────────────────────────────────────────────────
def style_run(run, pt=None, bold=None, color=None, font=FONT, italic=None):
    rPr = run._r.get_or_add_rPr()
    if pt is not None:
        rPr.set("sz", str(int(pt * 100)))
    if bold is not None:
        rPr.set("b", "1" if bold else "0")
    if italic is not None:
        rPr.set("i", "1" if italic else "0")
    for tag in ("a:latin", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
    if color is not None:
        run.font.color.rgb = color


def set_bullet(p, kind, level=0):
    """kind: None | 'num' | 'dot' | 'dash'."""
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if kind is None:
        pPr.set("marL", str(Emu(Inches(0.0))))
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
        return
    mar = 0.24 + 0.22 * level
    pPr.set("marL", str(int(Inches(mar))))
    pPr.set("indent", str(-int(Inches(0.22 if kind == "num" else 0.16))))
    if kind == "num":
        etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
    else:
        bf = etree.SubElement(pPr, qn("a:buFont"))
        bf.set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if kind == "dot" else "–")


def add_textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True, margin=0.04):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    return tb


def write_paras(tf, paras, *, first=True):
    """paras: list of dict(text|runs, pt, bold, color, bullet, level, after, align, font)."""
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if (first and i == 0) else tf.add_paragraph()
        runs = pd.get("runs") or [(pd["text"], {})]
        for txt, st in runs:
            r = p.add_run()
            r.text = txt
            style_run(r, pt=st.get("pt", pd.get("pt", 11)), bold=st.get("bold", pd.get("bold", False)),
                      color=st.get("color", pd.get("color", INK)), font=st.get("font", pd.get("font", FONT)))
        set_bullet(p, pd.get("bullet"), pd.get("level", 0))
        if pd.get("after") is not None:
            p.space_after = Pt(pd["after"])
        if pd.get("before") is not None:
            p.space_before = Pt(pd["before"])
        p.line_spacing = pd.get("spacing", 1.12)
        if pd.get("align"):
            p.alignment = pd["align"]


def paras_height(paras, width_in):
    h = 0.0
    for pd in paras:
        txt = pd.get("text") or "".join(t for t, _ in pd["runs"])
        pt = pd.get("pt", 11)
        ind = 0.0 if pd.get("bullet") is None else 0.24 + 0.22 * pd.get("level", 0)
        n = text_lines(txt, pt, width_in - ind - 0.08)
        h += n * line_h(pt, 1.2 * pd.get("spacing", 1.12) / 1.0) + (pd.get("after") or 0) / 72 + (pd.get("before") or 0) / 72
    return h + 0.1


def text_block(slide, x, y, w, paras, *, name="", h=None):
    need = paras_height(paras, w)
    tb = add_textbox(slide, x, y, w, h or need)
    write_paras(tb.text_frame, paras)
    if h is not None and need > h + 0.02:
        WARN.append(f"{name}: text needs {need:.2f}in > box {h:.2f}in")
    return y + (h or need)


def section_head(label, pt=12.5):
    return {"text": label, "pt": pt, "bold": True, "color": NAVY, "after": 3}


def body(text, bullet="dot", pt=11, color=INK, level=0, after=2, bold=False):
    return {"text": text, "pt": pt, "bullet": bullet, "color": color, "level": level, "after": after,
            "bold": bold}


def note(text, pt=9.5):
    return {"text": text, "pt": pt, "color": GRAY, "after": 2}


# ── shapes ─────────────────────────────────────────────────────────────────
def picture_fit(slide, path, x, y, w, h, *, align="center"):
    im = Image.open(path)
    ar = im.width / im.height
    if w / h > ar:
        pw, ph = h * ar, h
    else:
        pw, ph = w, w / ar
    px = x + (w - pw) / 2 if align == "center" else x
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(y), Inches(pw), Inches(ph))
    return px, y, pw, ph


def caption(slide, x, y, w, fig_no, text, pt=11):
    tb = add_textbox(slide, x, y, w, 0.34)
    write_paras(tb.text_frame, [{"runs": [(f"Fig {fig_no}. ", {"bold": True}), (text, {})],
                                 "pt": pt, "color": INK}])
    return y + 0.34


def table(slide, x, y, col_w, rows, *, pt=10, header=True, row_h=0.3, first_col_bold=False,
          header_fill=NAVY, zebra=True, col_align=None):
    n_r, n_c = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(col_w)),
                                   Inches(row_h * n_r))
    tbl = shape.table
    # 템플릿 기본 표 스타일 대신 직접 칠한다
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1" if header else "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    total_h = 0.0
    for i, row in enumerate(rows):
        max_lines = 1
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            is_head = header and i == 0
            runs = val if isinstance(val, list) else [(str(val), {})]
            p = tf.paragraphs[0]
            for txt, st in runs:
                r = p.add_run()
                r.text = txt
                style_run(r, pt=st.get("pt", pt), bold=st.get("bold", is_head or (first_col_bold and j == 0)),
                          color=st.get("color", RGBColor(0xFF, 0xFF, 0xFF) if is_head else INK))
            if col_align and col_align[j]:
                p.alignment = col_align[j]
            if is_head:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            elif zebra and i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            txt = "".join(t for t, _ in runs)
            max_lines = max(max_lines, text_lines(txt, pt, col_w[j] - 0.14))
        rh = max(row_h, max_lines * line_h(pt, 1.25) + 0.08)
        tbl.rows[i].height = Inches(rh)
        total_h += rh
    return y + total_h


def rounded(slide, x, y, w, h, fill=LIGHT, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def callout(slide, x, y, w, label, before, after, unit="", *, good=True, sep="  →  ",
            colors=None):
    """before → after 한 줄 요약 상자. sep='  vs  ' 면 개선이 아니라 나란히 비교."""
    rounded(slide, x, y, w, 0.95)
    tb = add_textbox(slide, x + 0.12, y + 0.06, w - 0.24, 0.3)
    write_paras(tb.text_frame, [{"text": label, "pt": 10.5, "color": GRAY}])
    tb = add_textbox(slide, x + 0.12, y + 0.36, w - 0.24, 0.55)
    c1, c2 = colors or (RED, GREEN if good else RED)
    write_paras(tb.text_frame, [{"runs": [(before, {"color": c1}), (sep, {"color": GRAY, "pt": 14}),
                                          (after, {"color": c2}),
                                          (unit, {"color": GRAY, "pt": 12})],
                                 "pt": 21, "bold": True}])


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def set_title(slide, text, font=FONT):
    ph = slide.shapes.title
    p = ph.text_frame.paragraphs[0]
    for r in list(p.runs):
        p._p.remove(r._r)
    r = p.add_run()
    r.text = text
    style_run(r, bold=True, font=font)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = next(i for i in ids if prs.part.related_part(i.rId) is slide.part)
    lst.remove(el)
    lst.insert(new_index, el)


# ── content ────────────────────────────────────────────────────────────────
SCRIPT: list[tuple[str, str]] = []   # (slide title, 대본) — md 와 notes 가 같은 것을 쓴다


def module_slide(slide, title, fig_no, fig_path, fig_caption, purpose, io_rows, method, *,
                 extra_fig=None, script=""):
    """모듈 슬라이드: 왼쪽 그림 · 오른쪽 목적 / Input·Output / 방법."""
    set_title(slide, title)
    for sh in list(slide.shapes):
        if not sh.is_placeholder:
            remove_shape(sh)
    # left
    LX, LW = 0.35, 6.15
    if extra_fig is None:
        im = Image.open(fig_path)
        fh = min(4.9, LW * im.height / im.width)
        top = 1.2 + max(0.0, (5.6 - fh - 0.42) / 2)
        _, fy, _, fh = picture_fit(slide, fig_path, LX, top, LW, fh)
        caption(slide, LX, fy + fh + 0.08, LW, fig_no, fig_caption)
    else:
        _, fy, _, fh = picture_fit(slide, fig_path, LX, 1.15, LW, 2.45)
        y2 = caption(slide, LX, fy + fh + 0.04, LW, fig_no, fig_caption)
        p2, no2, cap2 = extra_fig
        _, fy2, _, fh2 = picture_fit(slide, p2, LX, y2 + 0.12, LW, 2.75)
        caption(slide, LX, fy2 + fh2 + 0.04, LW, no2, cap2)
    # right
    RX, RW = 6.85, 6.15
    y = text_block(slide, RX, 1.1, RW, [section_head("1. 목적", pt=14)] +
                   [body(t, bullet=None, after=2, pt=12) for t in purpose], name=title + " 목적")
    y = text_block(slide, RX, y + 0.12, RW, [section_head("2. Input / Output", pt=14)], name=title)
    y = table(slide, RX + 0.05, y, [0.6, 2.3, 3.15],
              [["", "항목", "형태 · 의미"]] + io_rows, pt=10.5, row_h=0.31,
              col_align=[PP_ALIGN.CENTER, None, None])
    y = text_block(slide, RX, y + 0.2, RW, [section_head("3. 방법", pt=14)] +
                   [body(t, bullet="num", after=4, pt=12) for t in method], name=title + " 방법")
    if y > 7.1:
        WARN.append(f"{title}: right column ends at {y:.2f}in")
    set_notes(slide, script)
    SCRIPT.append((title, script))


def exp_slide(prs, layout, title, headline, fig, fig_no, fig_cap, blocks, *, fig_box=(0.35, 1.62, 7.35, 4.75),
              right_x=7.95, right_w=5.05, callouts=(), script=""):
    s = prs.slides.add_slide(layout)
    set_title(s, title)
    text_block(s, 0.35, 0.98, 12.6, [{"text": headline, "pt": 15, "bold": True, "color": NAVY}],
               name=title + " headline")
    left_bottom = None
    if fig is not None:
        x, y, w, h = fig_box
        _, fy, _, fh = picture_fit(s, fig, x, y, w, h)
        left_bottom = caption(s, x, fy + fh + 0.06, w, fig_no, fig_cap, pt=11)
    y = 1.62
    for head, items in blocks:
        paras = [section_head(head, pt=13)] + [
            body(t, pt=11.5, after=3) if isinstance(t, str) else t for t in items]
        y = text_block(s, right_x, y, right_w, paras, name=f"{title}/{head}") + 0.08
    if callouts and left_bottom is not None and left_bottom + 0.25 + 0.95 <= 7.0:
        # 그림 아래 빈 자리에 가로로 늘어놓는다
        x0, _, w0, _ = fig_box
        n = len(callouts)
        cw = (w0 - 0.2 * (n - 1)) / n
        cy = min(left_bottom + 0.3, 7.0 - 0.95)
        for i, c in enumerate(callouts):
            label, before, after, unit, *opt = c
            callout(s, x0 + i * (cw + 0.2), cy, cw, label, before, after, unit, **(opt[0] if opt else {}))
    else:
        for c in callouts:
            label, before, after, unit, *opt = c
            callout(s, right_x, y + 0.04, right_w, label, before, after, unit, **(opt[0] if opt else {}))
            y += 1.05
    if y > 7.15:
        WARN.append(f"{title}: right column ends at {y:.2f}in")
    set_notes(s, script)
    SCRIPT.append((title, script))
    return s


def main():
    prs = Presentation(str(TEMPLATE))
    S = list(prs.slides)
    layout = S[11].slide_layout  # '제목 및 내용'

    # 1 · 표지 — 날짜
    for sh in S[0].shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("2026."):
            runs = sh.text_frame.paragraphs[0].runs
            runs[0].text, runs[1].text, runs[2].text = "2026. 09. ", "30", "."
    s1 = ("안녕하십니까, 오늘은 Safety VLA 연구의 네 번째 세미나로, AG3S 파이프라인을 처음부터 끝까지 "
          "설명드리고, 그 과정에서 만난 기술적 어려움과 앞으로의 계획을 말씀드리겠습니다.")
    set_notes(S[0], s1)
    SCRIPT.append(("표지", s1))
    s2 = ("순서는 이렇습니다. 지난 발표 이후 받은 피드백을 먼저 짚고, AG3S 전체 구조를 한 장으로 보여 드린 뒤, "
          "다섯 모듈을 하나씩 설명드리겠습니다. 모듈 설명에서는 무엇을 왜 하는지에만 집중하고, "
          "구현하며 부딪힌 문제는 Experiment 에 모아서 말씀드리겠습니다.")
    set_notes(S[1], s2)
    SCRIPT.append(("0. Index", s2))
    SCRIPT.append(("1. 이전 연구 내용", "(발표자가 채운다)"))

    # 4 · 피드백 — 각 항목이 이번 발표 어디서 답해지는가
    map_lines = ["→  3. AG3S 아키텍처 (Fig 1)", "→  3. Perception 모듈이 왜 필요한가 (장·단점)",
                 "→  5. Experiment — 단계별 기술적 어려움", "→  5. Experiment — ①–⑧"]
    fb = [sh for sh in S[3].shapes if sh.has_text_frame and not sh.is_placeholder][0]
    ys = [1.48, 2.09, 2.70, 3.31]  # 템플릿 네 줄의 중심 1.68 · 2.29 · 2.90 · 3.51 in
    for yy, t in zip(ys, map_lines):
        tb = add_textbox(S[3], 8.45, yy, 4.6, 0.4, anchor=MSO_ANCHOR.MIDDLE)
        write_paras(tb.text_frame, [{"text": t, "pt": 13, "color": GRAY}])
    s4 = ("지난 발표 뒤 네 가지 피드백을 받았습니다. 전체 아키텍처와 데이터 흐름, Perception 모듈이 왜 필요한지, "
          "연구의 기술적 어려움, 그리고 단계별 연구 문제입니다. 오른쪽에 적은 것처럼, 앞의 둘은 3장 아키텍처에서, "
          "뒤의 둘은 5장 Experiment 에서 답하겠습니다.")
    set_notes(S[3], s4)
    SCRIPT.append(("2. 피드백", s4))

    # 5 · 아키텍처 — 그림은 사용자가 넣은 그대로 둔다
    s5 = ("전체 구조입니다. 한 문장으로 말씀드리면, VLA 의 visual attention 으로 작업 대상인 Target 과 장애물을 "
          "구분하고, Target 을 충돌 대상에서 뺀 거리장을 trajectory optimization 의 distance constraint 로 넘기는 "
          "구조입니다.\n\n"
          "왼쪽 local PC 에서 MuJoCo 가 observation 을 만들고, 오른쪽 GPU server 에서 π0.5, AG3S, trajopt 가 한 "
          "process 로 돕니다. 한 process 로 묶은 이유는 attention 이 model 내부 값이라 밖으로 꺼내지 않고 memory 로 "
          "넘기기 위해서입니다.\n\n"
          "π0.5 는 flow matching 으로 noise 에서 시작해 Euler step 10 번을 밟아 50 step × 16 차원 action chunk 를 "
          "만듭니다. 첫 step 인 k=0 에서 layer 8, head 2 의 attention, 즉 동작 token 이 이미지의 어디를 보는지를 "
          "읽어 AG3S 로 보냅니다.\n\n"
          "보라색 화살표를 따라가면 attention 이 3D 로 올라가고, Target 이 정해지고, cuRobo 가 Target 을 뺀 거리장을 "
          "만들고, 그것이 trajopt 의 거리 제약이 됩니다. 파란 화살표는 VLA 의 action chunk 가 최적화의 초기 궤적이자 "
          "기준이 되는 경로입니다. 결과는 local PC 로 돌아가 8 step 을 15 Hz 로 실행하고 다음 observation 을 찍습니다.\n\n"
          "cuRobo 와 trajopt 는 가져다 쓰는 외부 module 이라 점선 상자로 구분했습니다.")
    set_notes(S[4], s5)
    SCRIPT.append(("3. AG3S — 아키텍처", s5))

    # 5b · Perception 필요성 (피드백 2) — 새 슬라이드, 아키텍처 바로 뒤
    pn = prs.slides.add_slide(layout)
    set_title(pn, "3. AG3S – 왜 Perception 모듈이 필요한가")
    text_block(pn, 0.35, 0.98, 12.6, [{"text": "VLA 는 '무엇을 할지' 는 알지만 '어디에 무엇이 있는지' 를 기하로 보장하지 않는다",
                                         "pt": 15, "bold": True, "color": NAVY}], name="perception headline")
    cols = [
        ("VLA 단독의 한계", RED, [
            "action chunk 는 학습 분포를 따를 뿐, 충돌 여부를 계산하지 않는다",
            "학습에 없던 장애물 · 배치에서는 안전을 보장할 근거가 없다",
            "궤적을 고치려면 '무엇이 어디에 얼마나 떨어져 있는가' 가 필요하다"]),
        ("AG3S 의 장점", GREEN, [
            "VLA 를 다시 학습하지 않는다 — 정책의 attention 을 그대로 읽는다",
            "Target 은 잡게 두고 Obstacle 만 피하게 한다 (target-free layer)",
            "거리 제약 d − r − margin ≥ 0 으로 안전을 수치로 판정한다",
            "tracking cost 로 VLA 의 의도는 최대한 유지한다"]),
        ("AG3S 의 단점 · 비용", GRAY, [
            "depth · camera calibration · 로봇 모델 정확도에 의존한다",
            "attention 이 과제 국면에 따라 목적지로 옮겨간다",
            "지각 한 바퀴가 느리다 — 현재 chunk 예산을 넘는다 (5장)",
            "모듈이 늘어 실패 지점도 늘어난다 → 단계별 검증 필요"]),
    ]
    cx = 0.35
    for head, color, items in cols:
        rounded(pn, cx, 1.62, 4.1, 3.55, fill=LIGHT)
        text_block(pn, cx + 0.15, 1.75, 3.8, [{"text": head, "pt": 16, "bold": True, "color": color, "after": 8}] +
                   [body(t, pt=12.5, after=7) for t in items], name=f"perception/{head}", h=3.35)
        cx += 4.25
    text_block(pn, 0.35, 5.4, 12.6, [
        {"runs": [("정리 — ", {"bold": True, "color": NAVY}),
                  ("VLA 가 '무엇을' 정하고, AG3S 가 '무엇이 Target 인지' 와 '어디에 무엇이 있는지' 를 거리장으로 만들고, "
                   "trajopt 가 VLA 의 동작을 출발점으로 '어떻게 안전하게' 움직일지 고친다.", {})], "pt": 13}],
        name="perception summary")
    sp = ("피드백 두 번째, Perception 모듈이 왜 필요한가에 대한 답입니다.\n\n"
          "VLA 는 무엇을 할지는 잘 알지만, 내놓는 action chunk 가 충돌하는지는 계산하지 않습니다. 학습 분포를 따를 뿐이라 "
          "처음 보는 장애물이나 배치에서는 안전을 보장할 근거가 없습니다. 궤적을 고치려면 무엇이 어디에 얼마나 떨어져 "
          "있는지가 필요하고, 그것을 만드는 것이 Perception 모듈, AG3S 입니다.\n\n"
          "장점은 VLA 를 다시 학습하지 않고 attention 을 그대로 읽는다는 것, 그리고 Target 은 잡게 두고 장애물만 피하게 "
          "할 수 있다는 것입니다. 안전은 거리 제약으로 수치 판정하고, tracking cost 로 VLA 의 의도는 최대한 유지합니다.\n\n"
          "단점도 분명합니다. depth 와 calibration, 로봇 모델 정확도에 의존하고, attention 이 과제 국면에 따라 목적지로 "
          "옮겨가며, 지각 한 바퀴가 아직 느립니다. 모듈이 늘어난 만큼 실패 지점도 늘어서 단계별 검증이 필요했고, "
          "그 내용을 5장에서 말씀드리겠습니다.")
    set_notes(pn, sp)
    SCRIPT.append(("3. AG3S – 왜 Perception 모듈이 필요한가", sp))
    move_slide(prs, pn, 5)

    # 6 · Robot mask (1/5)
    module_slide(
        S[5], "4. AG3S – Robot Mask (1/5)", 2, FIG / "rm-scene.png",
        "세 camera 의 robot mask — 빨강 = 로봇 몸으로 판정해 지운 pixel",
        ["depth 에서 로봇 자기 몸이 찍힌 pixel 을 지워, 로봇이 자기 팔을 장애물로 보지 않게 한다. "
         "학습된 segmentation 이 아니라 로봇이 이미 아는 자기 자세(FK)로 지운다."],
        [["In", "depth × 3 · K · T_base←cam", "480×640 [m], head · 왼손목 · 오른손목"],
         ["In", "robot state q (16D)", "camera 마다 자기 촬영 시각의 자세"],
         ["In", "URDF sphere model", "로봇 몸을 구 여러 개로 근사"],
         ["Out", "masked depth × 3", "로봇 pixel = 0 → cuRobo TSDF 입력"],
         ["Out", "filtered point cloud", "pixel 위치 (u,v) 보존 → Attention lifting 입력"]],
        ["촬영 시각의 q 로 FK → 로봇 구 중심 위치 계산",
         "depth 를 3D 로 back-projection (각 점의 pixel 위치 u,v 보존)",
         "구 안에 든 점 = 로봇 → point cloud 에서 삭제",
         "삭제된 점의 (u,v) 를 image mask 로 되돌려 depth 에서 0 처리"],
        script=("첫 번째 모듈, Robot mask 입니다. 목적은 간단합니다. camera depth 에는 로봇 자기 팔도 찍히는데, 이것을 "
                "지우지 않으면 로봇이 자기 팔을 장애물로 보고 피하려 합니다.\n\n"
                "입력은 세 camera 의 depth 와 camera 파라미터, 그리고 로봇 자세입니다. 여기서 자세는 현재 값이 아니라 "
                "그 camera 가 찍힌 순간의 자세를 씁니다. 손목 camera 는 움직이기 때문에 시각이 어긋나면 엉뚱한 곳을 "
                "지우게 됩니다.\n\n"
                "방법은 학습된 segmentation 이 아니라, 로봇이 이미 아는 자기 자세를 이용합니다. forward kinematics 로 "
                "로봇을 근사한 구들의 위치를 구하고, depth 를 3D 점으로 올린 뒤 구 안에 든 점을 로봇으로 판정해 지웁니다. "
                "지운 점의 pixel 위치로 depth 이미지에도 mask 를 만듭니다.\n\n"
                "그림에서 빨간 부분이 지운 pixel 입니다. 머리 camera 에서 17 %, 왼손목 11 %, 오른손목 6 % 정도가 로봇입니다. "
                "출력은 두 갈래로, mask 된 depth 는 cuRobo 로, 점들은 Attention lifting 으로 갑니다."))

    # 7 · Robot mask on/off
    s7 = S[6]
    set_title(s7, "4. AG3S – Robot Mask (1/5)")
    for sh in list(s7.shapes):
        if not sh.is_placeholder:
            remove_shape(sh)
    _, fy, _, fh = picture_fit(s7, FIG / "rm-onoff.png", 0.35, 1.15, 7.6, 5.4)
    caption(s7, 0.35, fy + fh + 0.06, 7.6, 3, "같은 frame 에서 robot mask OFF / ON — ESDF 단면과 로봇 구 여유거리", pt=11)
    y = text_block(s7, 8.25, 1.15, 4.75, [
        section_head("Robot mask on / off 비교"),
        body("같은 frame · 같은 로봇 구 · 같은 cuRobo 설정에서 mask 만 끄고 켰다", pt=11, after=6),
        body("OFF — 팔이 찍힌 그대로 거리장에 적분된다. 로봇 구가 '자기 팔' 과 겹친 것으로 판정된다", pt=11, after=6),
        body("ON — 팔이 지워지고 테이블 · 물체만 거리장에 남는다", pt=11, after=6),
        body("여유거리 = d − r − margin  (0 이상이면 안전)", pt=11, after=6),
    ], name="rm onoff")
    callout(s7, 8.25, y + 0.1, 4.75, "위반 로봇 구 (120 개 중)", "105", "18", "  / 120")
    callout(s7, 8.25, y + 1.15, 4.75, "최악 여유거리", "−138.4", "−41.7", "  mm")
    text_block(s7, 8.25, y + 2.25, 4.75, [note("남은 음수는 손이 실제 물체(사과 · 바구니)에 가까이 간 구간이다 — "
                                               "이것을 고치는 것이 trajopt 의 몫이다")], name="rm onoff note")
    s7n = ("robot mask 를 끄고 켠 비교입니다. 같은 frame, 같은 로봇 구, 같은 cuRobo 설정에서 mask 만 바꿨습니다.\n\n"
           "위 두 그림은 camera 영상이 아니라 로봇 앞쪽을 세로로 자른 거리장 단면입니다. 왼쪽 mask 를 끈 경우 팔이 그대로 "
           "표면으로 적분돼서, 로봇 구들이 자기 팔과 겹친 것으로 나옵니다. 오른쪽은 팔이 지워지고 테이블과 물체만 남습니다.\n\n"
           "아래 그래프가 로봇 구 120 개의 여유거리, 즉 거리에서 구 반지름과 margin 을 뺀 값입니다. mask 가 없으면 120 개 중 "
           "105 개가 위반이고, mask 를 켜면 18 개로 줄어듭니다. 최악 여유거리도 −138 에서 −42 mm 로 올라옵니다. 남은 음수는 "
           "손이 실제로 사과와 바구니에 가까이 간 구간이고, 이것을 고치는 것이 뒤에 나올 trajopt 의 몫입니다.")
    set_notes(s7, s7n)
    SCRIPT.append(("4. AG3S – Robot Mask on/off", s7n))

    # 8 · Attention lifting
    module_slide(
        S[7], "4. AG3S – Attention lifting (2/5)", 4, FIG / "ap-attention.png",
        "Attention 16×16 grid → RGB overlay → 3D point (사과 접근 중)",
        ["VLA 가 이미지의 '어디를 보는가' (camera 별 16×16 attention) 를 3D 점마다 붙여, 3D 공간에서 Target 후보를 "
         "찾을 수 있게 한다. attention 이 낮다고 점을 버리지 않는다 — 충돌 기하는 3D 가 정한다."],
        [["In", "attention 16×16 × 3 camera", "π0.5 Euler step 0 · layer 8 · head 2"],
         ["In", "filtered point cloud", "Robot mask 출력, pixel 위치 (u,v) 보존"],
         ["In", "image size", "480×640 — attention 을 늘릴 해상도"],
         ["Out", "AttentionPointCloud", "점마다 attention 값 (정규화 0–1 · 원래 값)"]],
        ["16×16 attention → bilinear upsampling → 480×640 map",
         "각 3D 점이 자기 pixel (u,v) 의 attention 값을 읽는다",
         "세 camera 결과를 로봇 base 좌표계에서 voxel 단위로 합친다 (최댓값)",
         "percentile 정규화 — 크기만 맞추고 순위는 바꾸지 않는다"],
        script=("두 번째, Attention lifting 입니다. π0.5 가 주는 attention 은 camera 마다 16×16 격자, 즉 2D 입니다. "
                "그런데 충돌은 3D 에서 일어나므로, 이 2D 신호를 3D 점에 붙여야 합니다.\n\n"
                "그림 왼쪽이 세 camera 의 16×16 attention 이고, 가운데는 그것을 정책이 본 RGB 에 겹친 것입니다. 사과 위에 "
                "attention 이 모여 있습니다. 오른쪽이 결과로, 3D 점마다 attention 을 칠했더니 사과 자리가 밝게 나옵니다.\n\n"
                "방법은 16×16 을 이미지 해상도로 늘린 뒤, 각 3D 점이 자기가 찍힌 pixel 위치의 값을 읽는 것입니다. 세 camera "
                "결과는 로봇 좌표계에서 합칩니다.\n\n"
                "중요한 원칙은 attention 이 낮다고 점을 버리지 않는다는 것입니다. attention 은 어디가 Target 후보인지 "
                "알려 주는 신호일 뿐이고, 무엇이 장애물인지는 3D 기하가 정합니다."))

    # 9 · Target grounding
    module_slide(
        S[8], "4. AG3S – Target Grounding (3/5)", 5, FIG / "tg-method.png",
        "① attention 상위 점에서 출발 → ② 3D region growing → ③ 점수 (apple 이 1위)",
        ["attention 이 붙은 3D 점에서 '정책이 가리키는 물체' 하나를 Target 으로 고르고, 나머지 물체는 Obstacle 로 "
         "남긴다. attention 은 어디서 시작할지만 정하고, 물체의 경계는 3D 기하가 정한다."],
        [["In", "AttentionPointCloud", "점마다 attention 값"],
         ["In", "support surface", "테이블 · 바닥 평면 (RANSAC) — 연결에서만 제외"],
         ["Out", "Target points · centroid", "충돌 대상에서 뺄 물체"],
         ["Out", "후보 cluster 전체 · 상태", "Target 이 아닌 것 = Obstacle"]],
        ["Seed: attention 상위 5 % 점",
         "3D region growing — 붙어 있는 점을 한 물체(cluster)로 묶는다",
         "점수 = attention 크기 · 조밀도 · attention 최고점까지 거리 → 1위 = Target",
         "TargetConfirm — 다른 후보가 3 frame 연속 1위여야 Target 을 바꾼다"],
        script=("세 번째, Target grounding 입니다. attention 이 붙은 점들 중에서 정책이 가리키는 물체 하나를 Target 으로 "
                "고르고, 나머지는 장애물로 남기는 단계입니다.\n\n"
                "그림 왼쪽처럼 attention 상위 5 % 점에서 출발합니다. 가운데는 거기서 3D 로 붙어 있는 점들을 묶어 물체 "
                "단위 덩어리를 만든 것이고, 이 frame 에서는 7 개가 나왔습니다. 핵심은 attention 은 출발점만 정하고, 무엇이 "
                "한 물체인지는 기하가 정한다는 것입니다.\n\n"
                "오른쪽이 점수입니다. attention 크기와 점의 조밀도, attention 최고점까지 거리를 합쳐 순위를 매기고, 사과가 "
                "2 위보다 4.6 배 높은 점수로 1 위가 됐습니다.\n\n"
                "마지막으로 Target 이 frame 마다 흔들리지 않도록, 다른 후보가 세 frame 연속 1 위가 되어야 바꿉니다. 테이블 "
                "같은 지지면은 지우지 않고, 물체들을 잇지 않도록 연결에서만 뺍니다."))

    # 10 · TSDF → ESDF
    module_slide(
        S[9], "4. AG3S – TSDF → ESDF (cuRobo) (4/5)", 6, FIG / "esdf-role.png",
        "cuRobo 입력(mask 된 depth) → 3D point map → ESDF 표면",
        ["로봇이 지워진 depth 로 공간의 모든 칸(voxel)에 '가장 가까운 표면까지의 거리' 를 채운다. Target 을 뺀 "
         "target-free layer 를 따로 만들어, Target 을 충돌 대상에서 제외한 환경을 만든다."],
        [["In", "masked depth × 3 · K · T", "Robot mask 출력"],
         ["In", "Target points", "Target grounding 출력"],
         ["Out", "coarse ESDF", "20 mm voxel · 128³ (2.56 m) — 작업공간 전체"],
         ["Out", "fine ESDF", "5 mm voxel · 128³ (0.64 m) — 손 주변"],
         ["Out", "target-free fine layer", "Target 표면만 뺀 거리장"],
         ["Out", "label layer", "voxel 마다 가장 가까운 물체가 무엇인지"]],
        ["cuRobo Mapper — depth 를 적분해 TSDF (표면 근처만 거리 기록)",
         "같은 TSDF 에서 해상도만 바꿔 ESDF (모든 칸에 거리) 두 번 추출",
         "Target 표면 voxel 을 빼고 다시 전파 → target-free layer",
         "결과는 격자 값 · 원점 · voxel 크기로 넘긴다 (질의에 cuRobo 불필요)"],
        extra_fig=(FIG / "esdf-target-free.png", 7, "Target 높이의 수평 단면 — fine ESDF vs target-free layer (16D)"),
        script=("네 번째, cuRobo 로 거리장을 만드는 단계입니다. 여기서부터는 외부 module 인 cuRobo 를 씁니다.\n\n"
                "용어를 먼저 말씀드리면, TSDF 는 depth 를 쌓아 표면 근처 얇은 띠에만 거리를 적은 격자이고, ESDF 는 그것을 "
                "펼쳐 모든 칸에 가장 가까운 표면까지의 거리를 채운 격자입니다. 최적화가 실제로 묻는 것은 ESDF 입니다.\n\n"
                "위 그림이 흐름입니다. mask 된 depth 가 들어가고, 3D 점 지도를 거쳐, 오른쪽처럼 표면이 격자로 표현됩니다. "
                "같은 TSDF 에서 해상도만 바꿔 두 번 뽑는데, 작업공간 전체를 덮는 20 mm 격자와 손 주변만 덮는 5 mm 격자입니다. "
                "먼 곳은 거친 격자가, 가까운 곳은 정밀한 격자가 맡습니다.\n\n"
                "이 연구에서 중요한 것은 아래 그림입니다. Target 의 표면만 빼고 거리를 다시 전파한 target-free layer 를 "
                "따로 만듭니다. 사과 높이에서 자른 단면인데, 왼쪽에는 빨간 원 안에 사과가 있고 오른쪽에서는 사과만 사라지고 "
                "다른 물체는 그대로 남아 있습니다. 이것이 Target 을 충돌 대상에서 제외한다는 것의 실체입니다."))

    # 11 · Distance field adapter
    module_slide(
        S[10], "4. AG3S – Distance Field Adapter (5/5)", 8, FIG / "adapter-top.png",
        "(a) 단면 위 로봇 구와 ∇d  (b) 로봇 구 여유거리  (c) 제약 행 d − r − margin",
        ["cuRobo 의 여러 거리 격자를 trajopt 가 묻는 두 함수 distance(p) · gradient(p) 로 바꾸고, 물체 종류에 따라 "
         "요구하는 여유(margin)를 다르게 준다."],
        [["In", "ESDF layers + label", "coarse · fine · target-free"],
         ["In", "query point p(q)", "로봇 구 중심 약 120 개"],
         ["Out", "distance(p)  (N,)", "trilinear 보간, 음수 = 표면 안쪽"],
         ["Out", "gradient(p)  (N,3)", "표면에서 멀어지는 방향"],
         ["Out", "margin by label", "obstacle 50 mm · destination 20 mm"]],
        ["층별 trilinear 보간 → d = min(coarse, fine)",
         "gradient 는 d 를 낸 바로 그 층에서 계산 — 값과 방향이 어긋나지 않게",
         "로봇 구 하나당 제약 한 행: d(p(q)) − r − margin ≥ 0",
         "Target 을 잡아도 되는 링크(손가락)만 target-free layer 로 질의"],
        script=("다섯 번째, Distance field adapter 입니다. cuRobo 가 만든 것은 숫자가 채워진 격자일 뿐이고, 최적화가 "
                "필요한 것은 '이 점에서 거리가 얼마이고, 어느 방향으로 가면 멀어지는가' 두 가지입니다. 이 둘을 함수로 "
                "바꿔 주는 얇은 연결층입니다.\n\n"
                "로봇을 구 약 120 개로 근사해 두고, 각 구 중심을 물으면 거리와 기울기를 돌려줍니다. 거친 격자와 정밀한 "
                "격자 중 작은 값을 쓰고, 기울기는 그 값을 낸 같은 층에서 계산해야 최적화가 수렴합니다.\n\n"
                "그림 (a) 의 초록 화살표가 각 구가 받은 기울기, 즉 표면에서 멀어지는 방향입니다. (b) 는 구마다 여유거리를 "
                "칠한 것이고, (c) 가 최적화에 들어가는 제약 전부입니다. 구 하나당 '거리 − 반지름 − margin 이 0 이상' 이라는 "
                "한 행입니다.\n\n"
                "여기서 물체 종류에 따라 margin 이 달라집니다. 일반 장애물은 50 mm, 물건을 넣을 목적지는 20 mm 이고, Target 을 "
                "잡아야 하는 손가락만 target-free layer 를 물어서 사과에 닿을 수 있습니다."))

    # 12 · Experiment — 환경 (기존 슬라이드 재사용)
    e0 = S[11]
    set_title(e0, "5. Experiment – 실험 환경과 검증 단계")
    text_block(e0, 0.35, 0.98, 12.6, [{"text": "모듈이 옳은가 → 한 번의 closed loop 이 서는가 — 세 단계로 나눠 검증했다",
                                         "pt": 15, "bold": True, "color": NAVY}], name="env headline")
    table(e0, 0.35, 1.65, [1.55, 4.6],
          [["항목", "설정"],
           ["Robot · Sim", "Rainbow RB-Y1 양팔 로봇 · MuJoCo"],
           ["VLA", "π0.5, 16D LoRA fine-tune (30k step)"],
           ["Task", "\"Pick up the apple and place it in the basket.\""],
           ["Episode", "held-out test split 1800–1999 (학습에 안 쓴 씬)"],
           ["Control", "15 Hz · chunk 50 step 중 8 step 실행 (533 ms)"],
           ["Server", "π0.5 + AG3S + trajopt 한 process (GPU server)"]],
          pt=12, row_h=0.46, first_col_bold=True)
    stages = [("Offline replay", "저장된 기록을 다시 돌려\n모듈 출력을 한 frame 씩 검사"),
              ("Shadow", "실시간으로 판정만 하고\n로봇은 VLA 원본을 실행"),
              ("Closed loop", "trajopt 결과를\n실제로 실행")]
    x = 6.85
    for i, (h, d) in enumerate(stages):
        rounded(e0, x, 1.65, 1.85, 1.75, fill=LIGHT)
        text_block(e0, x + 0.1, 1.75, 1.65, [{"text": f"{i + 1}. {h}", "pt": 13, "bold": True, "color": NAVY,
                                                "after": 4}, {"text": d, "pt": 11.5}], name=f"stage {h}", h=1.5)
        if i < 2:
            arr = e0.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.88), Inches(2.4), Inches(0.24), Inches(0.28))
            arr.fill.solid()
            arr.fill.fore_color.rgb = NAVY
            arr.line.fill.background()
        x += 2.13
    text_block(e0, 6.85, 3.7, 6.15, [
        section_head("평가 지표", pt=14),
        body("여유거리 (clearance) = d − r − margin — 0 이상이면 안전", pt=12, after=5),
        body("feasible / violated — 최적화 뒤 최악 여유거리가 0 이상인가 (이진 판정)", pt=12, after=5),
        body("grasp · place 성공 — MuJoCo 참값으로 사과가 들리고 바구니 안에 놓였는가", pt=12, after=5),
        body("지연 — chunk 한 번에 드는 시간 vs 예산 533 ms", pt=12, after=5),
    ], name="metrics")
    se0 = ("실험 환경입니다. Rainbow RB-Y1 양팔 로봇을 MuJoCo 에서 돌리고, π0.5 를 16 차원 action 으로 LoRA fine-tune 한 "
           "model 을 씁니다. 과제는 사과를 집어 바구니에 넣는 것이고, 평가는 학습에 쓰지 않은 test split 에피소드에서 합니다. "
           "15 Hz 로 제어하고, 50 step chunk 중 8 step 만 실행한 뒤 다시 계획하므로 한 번의 계획에 주어진 시간은 533 ms 입니다.\n\n"
           "검증은 세 단계로 나눴습니다. 먼저 저장된 기록을 다시 돌려 모듈 출력을 frame 단위로 검사하고, 다음으로 실시간으로 "
           "판정만 하고 로봇은 VLA 원본을 실행하는 shadow, 마지막으로 최적화 결과를 실제로 실행하는 closed loop 입니다.\n\n"
           "지표는 여유거리와 그것을 0 기준으로 나눈 feasible / violated, 그리고 과제 성공, 지연입니다.")
    set_notes(e0, se0)
    SCRIPT.append(("5. Experiment – 실험 환경", se0))

    # 13 · 요약표
    sm = prs.slides.add_slide(layout)
    set_title(sm, "5. Experiment – 단계별 기술적 어려움 요약")
    text_block(sm, 0.35, 0.98, 12.6, [{"text": "구현이 아니라 '왜 어려운가' 로 정리한 여덟 가지 — 각각 다음 슬라이드에서 설명",
                                         "pt": 15, "bold": True, "color": NAVY}], name="summary headline")
    table(sm, 0.35, 1.6, [0.45, 1.75, 4.95, 3.35, 2.1],
          [["#", "단계", "어려움", "핵심 수치", "상태"],
           ["①", "Robot mask", "손이 Target 에 닿는 순간 self-filter 가 Target 까지 지운다",
            "사과 통째 삭제 17 → 0 frame", "해결 (검증 대기)"],
           ["②", "Target grounding", "주목하는 물체 ≠ 조작하는 물체 — 파지 후 attention 이 바구니로 간다",
            "Target 이름 뒤집힘 9 → 3 회", "해결 · 일부 남음"],
           ["③", "TSDF → ESDF", "cuRobo 는 틀린 입력에도 그럴듯한 거리장을 낸다", "|∇d| 0.25 → 1.00", "해결"],
           ["④", "Distance field", "20 mm 격자는 위험한 쪽(더 멀다)으로 틀린다", "최악 오차 +7.56 → −0.49 mm", "해결"],
           ["⑤", "Margin 정책", "margin 하나로 '피하기' 와 '넣기' 를 동시에 못 한다", "담기 여유 −5.3 → +24.7 mm", "해결"],
           ["⑥", "Closed loop", "각 단계는 정상인데 closed loop 에서 grasp 실패", "사과 들림 245.2 vs 0.0 mm", "원인 규명 · 수정"],
           ["⑦", "Closed loop", "거리장을 켜면 성공률이 떨어진다", "E3 성공 0/6 → 2/6", "진행 중"],
           ["⑧", "System", "camera 시차와 실시간 예산", "P50 2519 ms vs 533 ms", "시차 해결 · 속도 보류"]],
          pt=12, row_h=0.52, col_align=[PP_ALIGN.CENTER, None, None, None, None])
    ssm = ("구현하면서 만난 문제를 '무엇을 고쳤나' 가 아니라 '왜 어려운가' 로 다시 정리하면 이 여덟 가지입니다.\n\n"
           "지각 쪽에서는 로봇 mask 가 손이 닿는 순간 Target 까지 지우는 문제, attention 이 가리키는 물체와 실제로 쥔 물체가 "
           "달라지는 문제가 있었습니다. 거리장 쪽에서는 cuRobo 가 조용히 틀리는 문제, 거친 격자가 위험한 방향으로 틀리는 문제, "
           "margin 하나로는 피하기와 넣기를 동시에 못 하는 문제가 있었습니다.\n\n"
           "마지막 셋은 시스템 수준입니다. 단계마다는 정상인데 closed loop 에서 잡기가 실패하는 문제, 거리장을 켜면 성공률이 "
           "떨어지는 문제, 그리고 camera 시차와 실시간 예산입니다. 하나씩 말씀드리겠습니다.")
    set_notes(sm, ssm)
    SCRIPT.append(("5. Experiment – 요약", ssm))

    ex = []
    ex.append(exp_slide(
        prs, layout, "5. Experiment – ① Robot mask",
        "손이 사과에 닿는 순간, 로봇을 지우는 규칙이 사과까지 지운다",
        FIG / "x-selffilter.png", 9, "사과는 camera 에 보이지만(좌: 유지) 손이 도착하면 로봇 구 안에 들어 전부 지워진다(중·우)",
        [("왜 어려운가", ["self-filter 는 '로봇 구 안의 점 = 로봇' 이라는 기하 규칙이다",
                         "grasp 는 정의상 손이 Target 표면에 닿는 동작 — 로봇과 Target 이 같은 자리에 온다",
                         "덜 지우면 로봇이 장애물로 남고, 더 지우면 Target 이 사라진다 (trade-off)"]),
         ("측정", ["사과가 통째로 지워진 구간 17 frame — 원인은 손가락이 아니라 손목 링크의 큰 capsule (반지름 75 mm · 길이 250 mm)",
                  "모델 손가락이 늘 닫힌 자세로 들어가 실제보다 51 mm 안쪽에 있었다"]),
         ("해결", ["self-filter 전용 로봇 모델을 mesh 표면에 맞춘 덮개로 교체 · 손가락은 측정 개도로"])],
        callouts=[("사과 통째 삭제 frame (자가 점검)", "17", "0", "  frame"),
                  ("사과 점 보존율 (자가 점검)", "0.47", "0.92", "")],
        script=("첫 번째 어려움은 Robot mask 에서 나왔습니다. 그림 왼쪽은 손이 오기 전이라 사과가 초록, 즉 유지되지만, 손이 도착한 "
                "가운데와 오른쪽에서는 사과가 빨강, 전부 지워집니다. 사과는 camera 에 분명히 보이는데 말입니다.\n\n"
                "왜 어려운가 하면, self-filter 는 로봇 구 안의 점을 로봇으로 보는 규칙인데, grasp 는 정의상 손이 Target 표면에 "
                "닿는 동작입니다. 덜 지우면 로봇이 장애물로 남고, 더 지우면 Target 이 사라지는 trade-off 가 생깁니다.\n\n"
                "측정해 보니 사과가 통째로 지워진 17 frame 의 원인은 손가락이 아니라 손목 링크를 감싼 큰 capsule 이었고, 모델의 "
                "손가락이 늘 닫힌 자세로 들어가 실제보다 51 mm 안쪽에 있었습니다.\n\n"
                "self-filter 전용 로봇 모델을 mesh 표면에 맞춰 다시 만들고 손가락을 실제 개도로 놓았더니, 구현자 자가 점검 기준으로 "
                "통째 삭제가 17 frame 에서 0 으로, 사과 점 보존율이 0.47 에서 0.92 로 올랐습니다. 공식 검증은 남아 있습니다.")))
    ex.append(exp_slide(
        prs, layout, "5. Experiment – ② Target grounding",
        "정책이 '보는' 물체와 로봇이 '쥔' 물체가 다르다",
        FIG / "x-target-drift.png", 10,
        "ep1807 한 에피소드 — 1위 cluster · attention peak · 과제 국면 (위), Target 점수 (아래)",
        [("왜 어려운가", ["VLA 의 attention 은 과제 국면에 따라 사과 → 목적지(바구니)로 옮겨간다",
                         "충돌 제외와 접촉 허용은 '쥔 물체' 에 붙어야 한다",
                         "손이 사과를 가리면 사과 점이 사라져, 남은 물체가 1위가 된다"]),
         ("측정", ["seq 16 에 1위가 사과 → 바구니로 바뀜. 그러나 attention peak 가 바구니로 간 것은 seq 23",
                  "→ attention 이 옮겨서가 아니라 사과 점이 사라져 '남은 유일한 후보' 가 1위가 된 것 (점수 0.015)"]),
         ("해결 · 남은 문제", ["1위 선택 + 3 frame 확인(hysteresis) · 조작 대상 ID 유지 · 바구니는 조작 대상이 될 수 없게",
                             "남은 문제: 사과가 이웃 과일과 한 덩어리로 묶이는 경우 (ep1800 · ep1808)"])],
        callouts=[("연속 chunk 사이 Target 이름 뒤집힘", "9", "3", "  회"),
                  ("75 chunk 중 safe 판정", "14", "68", "  / 75")],
        script=("두 번째는 Target grounding 입니다. 정책이 보는 물체와 로봇이 쥔 물체가 다르다는 문제입니다.\n\n"
                "그림은 한 에피소드 동안 1 위 덩어리, attention 최고점, 과제 국면을 시간축에 놓은 것입니다. seq 16 에 1 위가 사과에서 "
                "바구니로 바뀌는데, attention 최고점이 바구니로 넘어간 것은 seq 23 입니다. 즉 attention 이 옮겨서가 아니라, 손이 "
                "사과를 가려 사과 점이 사라지면서 바구니가 남은 유일한 후보라 1 위가 된 것입니다. 아래 점수가 거의 0 으로 떨어진 "
                "구간이 그것입니다.\n\n"
                "어려운 이유는 attention 은 원래 과제 국면에 따라 목적지로 옮겨가는 것이 정상인데, 충돌 제외와 접촉 허용은 쥔 "
                "물체에 붙어야 한다는 데 있습니다.\n\n"
                "1 위 선택에 3 frame 확인을 더하고, 조작 대상 ID 를 유지하고, 바구니는 조작 대상이 될 수 없게 했습니다. 그 결과 "
                "Target 이름 뒤집힘이 9 회에서 3 회로, safe 판정이 75 chunk 중 14 에서 68 로 늘었습니다. 사과가 이웃 과일과 한 "
                "덩어리로 묶이는 경우는 아직 남아 있습니다.")))
    ex.append(exp_slide(
        prs, layout, "5. Experiment – ③ TSDF → ESDF",
        "cuRobo 는 틀린 입력에도 '그럴듯한' 거리장을 낸다 — 조용한 실패를 스스로 잡아야 한다",
        FIG / "x-eikonal.png", 11, "자기 진단 — 자유공간에서 |∇d| 분포 (거친 20 mm · 미세 5 mm)",
        [("왜 어려운가", ["거리장은 틀려도 숫자가 나온다 — 오류가 나지 않는다",
                         "격자 원점 · 버퍼 재사용 같은 API 규약이 어긋나면 값이 조용히 틀린다"]),
         ("측정", ["compute_esdf() 결과를 복사하지 않으면 다음 호출이 덮어쓴다 → |∇d| = 0.25",
                  "esdf_origin 은 격자 코너가 아니라 중심 — 코너로 넣으면 두 계층이 130–257 mm 어긋난다"]),
         ("해결", ["eikonal 자기 진단: 제대로 된 거리장이면 어디서든 |∇d| ≈ 1 (1 m 가면 거리가 1 m 바뀐다)",
                  "이 검사를 통과하지 못하면 거리장을 쓰지 않는다"])],
        fig_box=(0.9, 1.62, 6.3, 4.75),
        callouts=[("자유공간 |∇d| 중앙값", "0.25", "1.00", "  (coarse 0.999)")],
        script=("세 번째는 cuRobo 로 거리장을 만드는 단계입니다. 거리장은 틀려도 숫자가 나오고, 오류 메시지가 없습니다. 그래서 "
                "조용한 실패를 스스로 잡는 장치가 필요했습니다.\n\n"
                "실제로 두 번 밟았습니다. 하나는 cuRobo 의 결과 버퍼를 복사하지 않으면 다음 호출이 덮어쓰는 문제였고, 다른 하나는 "
                "격자 원점이 코너가 아니라 중심이라는 규약이었습니다. 원점을 잘못 넣으면 두 격자가 13 에서 26 cm 까지 어긋납니다.\n\n"
                "이것을 잡은 것이 eikonal 자기 진단입니다. 제대로 된 거리장이면 1 m 움직일 때 거리도 1 m 바뀌므로 기울기 크기가 "
                "어디서든 1 입니다. 그림이 자유공간에서 잰 기울기 크기 분포인데, 두 격자 모두 1 에 모여 있습니다. 버퍼 문제가 "
                "있을 때는 이 값이 0.25 였습니다. 지금은 이 검사를 통과하지 못하면 거리장을 쓰지 않습니다.")))
    ex.append(exp_slide(
        prs, layout, "5. Experiment – ④ Distance field",
        "20 mm 격자는 '실제보다 멀다' 는 쪽, 즉 위험한 쪽으로 틀린다",
        FIG / "x-two-tier.png", 12, "거리장 − 참값 [mm] 단면 — (b) 거친 20 mm 단독 (c) 2계층 20 + 5 mm (빨강 = 실제보다 멀다고 답함)",
        [("왜 어려운가", ["격자가 거칠면 표면 위치를 칸 중심에 뭉뚱그려 거리가 어긋난다",
                         "판정을 좌우하는 최악 구에서 '더 멀다' 로 틀리면 충돌을 놓친다",
                         "그렇다고 5 mm 격자를 전역에 두면 비용이 크다"]),
         ("해결", ["2계층 — 전역 20 mm + 손 주변 5 mm, 질의는 둘 중 작은 값",
                  "미세 창을 Target 중심이 아니라 손이 지나갈 경로(execution path)에 둔다"]),
         ("결과", ["손 쪽 제약 행이 미세 창 안에 드는 비율 0.06–0.10 → 1.00"])],
        fig_box=(0.35, 1.62, 7.35, 3.6),
        callouts=[("판정 지점 최악 구 오차 (+ = 낙관)", "+7.56", "−0.49", "  mm")],
        script=("네 번째는 거리장의 정밀도입니다. 격자가 20 mm 로 거칠면 표면을 칸 중심에 뭉뚱그리기 때문에 거리가 어긋나는데, "
                "문제는 그 방향입니다. 그림 가운데 빨간 부분이 '실제보다 멀다' 고 답한 곳, 즉 위험한 쪽으로 틀린 곳입니다.\n\n"
                "판정을 좌우하는 최악의 로봇 구에서 거친 격자는 7.56 mm 낙관적이었습니다. 그렇다고 5 mm 격자를 작업공간 전체에 "
                "두면 비용이 너무 큽니다.\n\n"
                "그래서 두 층을 씁니다. 전역은 20 mm, 손 주변은 5 mm 로 덮고 둘 중 작은 값을 씁니다. 오른쪽 그림처럼 빨간 부분이 "
                "사라지고, 최악 구 오차가 −0.49 mm, 오히려 약간 보수적인 쪽이 됐습니다. 이후 미세 창을 Target 중심이 아니라 손이 "
                "지나갈 경로에 두도록 바꿨고, 손 쪽 제약이 미세 창 안에서 답을 받는 비율이 10 % 이하에서 100 % 가 됐습니다.")))
    ex.append(exp_slide(
        prs, layout, "5. Experiment – ⑤ Margin 정책",
        "margin 하나로는 '장애물은 피하고, 바구니 안에는 넣는' 두 가지를 동시에 못 한다",
        FIG / "x-dest-margin.png", 13, "쥔 사과의 여유거리 — 전역 margin 50 mm (빨강) vs 목적지 margin 20 mm (초록)",
        [("왜 어려운가", ["전역 50 mm 를 요구하면 사과를 바구니 안으로 넣는 동작 자체가 위반이 된다",
                         "margin 을 풀면 그 완화가 로봇 팔 같은 다른 제약으로 새면 안 된다",
                         "거리장은 '어느 물체의 표면인지' 모른다 — 익명의 거리만 준다"]),
         ("해결", ["label layer 로 가장 가까운 표면이 어느 물체인지 알아낸다",
                  "쥔 물체 ↔ 목적지만 20 mm, 로봇 구 ↔ 목적지 · 그 밖은 50 mm 유지",
                  "margin 20 mm = 관측 여유 32.9 mm − 거친 격자 낙관 7.56 mm 의 상한 안"]),
         ("확인", ["로봇 구 쪽 여유거리는 두 설정에서 완전히 같다 — 완화가 새지 않았다"])],
        fig_box=(0.9, 1.62, 6.3, 4.75),
        callouts=[("담는 순간 쥔 사과의 여유거리", "−5.3", "+24.7", "  mm")],
        script=("다섯 번째는 margin 정책입니다. 장애물에서 50 mm 떨어지라는 규칙 하나를 전체에 적용하면, 사과를 바구니 안에 넣는 "
                "동작 자체가 위반이 됩니다. 그림의 빨간 막대가 전역 50 mm 일 때 쥔 사과의 여유거리인데, 담는 순간인 마지막 frame 에서 "
                "−5.3 mm 로 위반입니다.\n\n"
                "어려운 점은 거리장이 어느 물체의 표면인지 모르는 익명의 거리만 준다는 것, 그리고 margin 을 풀어도 그 완화가 로봇 "
                "팔 같은 다른 제약으로 새면 안 된다는 것입니다.\n\n"
                "label layer 로 가장 가까운 표면이 어느 물체인지 알아내고, 쥔 물체와 목적지 사이만 20 mm 로 줄였습니다. 20 mm 는 "
                "관측된 여유 32.9 mm 에서 거친 격자의 낙관 7.56 mm 를 뺀 상한 안에서 정했습니다. 초록 막대처럼 +24.7 mm 로 담기가 "
                "풀렸고, 로봇 팔 쪽 여유거리는 두 설정에서 완전히 같아 완화가 새지 않았음을 확인했습니다.")))
    ex.append(exp_slide(
        prs, layout, "5. Experiment – ⑥ Closed loop",
        "각 단계는 모두 정상인데, closed loop 에서는 사과를 들지 못한다",
        FIG / "x-hold-vs-shadow.png", 14,
        "왼쪽 closed loop (충돌 제약 OFF) vs 오른쪽 shadow — 손가락–사과 거리 · gripper · 사과 들림",
        [("왜 어려운가", ["단계별 판정만으로는 closed loop 실패를 설명할 수 없다",
                         "AG3S · cuRobo · trajopt 를 하나씩 재도 전부 정상이었다"]),
         ("측정", ["refined 가 reference 보다 참값 여유를 나쁘게 만든 chunk 0/75",
                  "충돌 제약을 꺼서 궤적 수정이 0.00° 여도 실패 — 원인은 최적화가 아니다",
                  "안전 판정 → HOLD(정지) 경로가 gripper 닫힘 명령을 삼켰다"]),
         ("해결", ["HOLD 중에도 닫힘 명령 유지 · HOLD 중 팔 드리프트 105.1 → 6.0 mm",
                  "QP 가 거절한 후보를 돌려주던 경로 수정"])],
        fig_box=(0.35, 1.62, 7.35, 3.3),
        callouts=[("사과 들림 — closed loop  vs  shadow (같은 서버)", "0.0 mm", "245.2 mm", "",
                   {"sep": "   vs   ", "colors": (RED, BLUE)})],
        script=("여섯 번째부터는 시스템 수준의 문제입니다. 모듈을 하나씩 재면 전부 정상인데, closed loop 에서는 사과를 들지 "
                "못했습니다. 같은 서버에서 로봇이 VLA 원본을 실행하는 shadow 는 사과를 245 mm 들어 올렸는데, 최적화 결과를 실제로 "
                "실행하면 0 mm 였습니다.\n\n"
                "AG3S, cuRobo, trajopt 를 하나씩 재 봤습니다. 최적화가 참값 여유를 나쁘게 만든 chunk 는 75 개 중 0 이었습니다. "
                "그림 왼쪽은 충돌 제약을 아예 꺼서 궤적 수정이 0 도인데도 실패한 경우입니다. 원인은 최적화가 아니었습니다.\n\n"
                "범인은 안전 판정 뒤의 정지 경로였습니다. 위험하다고 판단하면 로봇을 HOLD 시키는데, 그 과정에서 gripper 닫힘 "
                "명령까지 삼켜서 손이 사과 옆에서 멈춰 있었던 것입니다. HOLD 중에도 닫힘 명령을 유지하고 팔이 흘러내리지 않게 "
                "고쳤습니다. 교훈은, 단계별로 결백해도 시스템은 실패할 수 있고, 그래서 추측이 아니라 측정으로 좁혀야 했다는 것입니다.")))

    # ⑦ 성공률 사다리 — native chart
    s = prs.slides.add_slide(layout)
    set_title(s, "5. Experiment – ⑦ Closed loop 성공률")
    text_block(s, 0.35, 0.98, 12.6, [{"text": "기능을 하나씩 켜며 성공률을 쟀다 — 거리장을 켜는 순간 무너지고, 원인을 고쳐 일부 회복",
                                        "pt": 15, "bold": True, "color": NAVY}], name="ladder headline")
    cd = CategoryChartData()
    cd.categories = ["E0 VLA 단독", "E1 +추적만", "E2 +목적함수·limit", "E3a +거리장", "E3b +거리장·판정"]
    cd.add_series("수정 전", (7 / 12 * 100, 6 / 12 * 100, 100.0, 0.0, 0.0))
    cd.add_series("수정 후", (None, None, None, 2 / 6 * 100, 2 / 6 * 100))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.35), Inches(1.6), Inches(7.6), Inches(4.8), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11)
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = 100, 0
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE3, 0xE3)
    va.tick_labels.font.size = Pt(10)
    va.tick_labels.font.color.rgb = GRAY
    va.has_title = True
    va.axis_title.text_frame.text = "성공률 [%]"
    va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(10.5)
    ca.format.line.color.rgb = LINE
    labels = [["7/12", "6/12", "6/6", "0/6", "0/6"], [None, None, None, "2/6", "2/6"]]
    for si, (ser, color) in enumerate(zip(ch.plots[0].series, [RGBColor(0x9A, 0xA4, 0xB8), NAVY])):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = color
        for pi, lab in enumerate(labels[si]):
            if lab is None:
                continue
            dl = ser.points[pi].data_label
            dl.has_text_frame = True
            dl.text_frame.text = lab
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
            dl.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.plots[0].gap_width = 60
    caption(s, 0.35, 6.45, 7.6, 15, "episode 1800 · 1807 · 1808 × 반복 — '수정 후' 는 여유 규칙 · camera 시각 · 거짓 attach 수정 뒤", pt=10)
    y = text_block(s, 8.25, 1.62, 4.75, [
        section_head("사다리의 의미", pt=13),
        body("E0 → E2: VLA 를 따라가기만 하는 최적화는 성공률을 떨어뜨리지 않는다", pt=11.5),
        body("E3: 거리장을 켜면 0/6 — 거리 제약이 잡기를 막았다", pt=11.5, after=6),
        section_head("E3 실패에서 찾은 원인 → 수정", pt=13),
        body("접근 단계의 20 mm 여유가 사과를 잡아도 되는 손가락에도 걸렸다 → 위반 159 → 13 chunk", pt=11.5),
        body("세 camera 의 촬영 시각이 달라 판정이 '미인증' 으로 멈췄다 → 관측 하나에 시각 하나", pt=11.5),
        body("빈손으로 닫아도 '쥐었다' 고 판정 → 거짓 attach 10/24 → 0/24", pt=11.5, after=6),
        section_head("남은 실패 (ep1800 · ep1808)", pt=13),
        body("충돌 위반은 0 인데 손목 관절이 5–8.7° 틀어진다", pt=11.5),
        body("최적화의 관절 한계(URDF)가 실제 로봇 범위보다 좁다 — 정책은 그 밖의 값을 낸다", pt=11.5),
    ], name="ladder text")
    sl = ("일곱 번째는 closed loop 성공률입니다. 기능을 하나씩 켜면서 쟀습니다. E0 는 VLA 단독, E1 은 최적화가 VLA 를 따라가기만 "
          "하는 것, E2 는 목적함수와 관절 한계까지 넣은 것, E3 는 거리장까지 켠 것입니다.\n\n"
          "E2 까지는 성공률이 떨어지지 않았습니다. 최적화 틀 자체가 VLA 의 동작을 망치지는 않는다는 뜻입니다. 그런데 거리장을 "
          "켜는 순간 6 번 중 0 번이 됐습니다.\n\n"
          "원인은 세 가지였습니다. 접근 단계의 20 mm 여유 규칙이 사과를 잡아도 되는 손가락에도 걸려 있었고, 세 camera 의 촬영 "
          "시각이 달라 판정이 멈췄고, 빈손으로 닫아도 쥐었다고 판정하는 경우가 있었습니다. 셋을 고친 뒤 E3 가 6 번 중 2 번으로 "
          "올라왔습니다.\n\n"
          "다만 성공은 아직 한 에피소드뿐입니다. 남은 실패에서는 충돌 위반이 없는데 손목 관절이 5 에서 9 도 틀어지는데, 최적화에 "
          "넣은 관절 한계가 실제 로봇 범위보다 좁은 것이 유력한 원인이라 지금 고치고 있습니다.")
    set_notes(s, sl)
    SCRIPT.append(("5. Experiment – ⑦ 성공률", sl))
    ex.append(s)

    # ⑧ timing — native bar chart
    s = prs.slides.add_slide(layout)
    set_title(s, "5. Experiment – ⑧ Camera 시차와 실시간 예산")
    text_block(s, 0.35, 0.98, 12.6, [{"text": "533 ms 안에 지각 · 거리장 · 최적화를 끝내야 하는데, 지금은 4.7 배 걸린다",
                                        "pt": 15, "bold": True, "color": NAVY}], name="timing headline")
    cd = CategoryChartData()
    cd.categories = ["chunk 예산", "TO (SQP + QP)", "AG3S (지각 한 바퀴)", "서버 total"]
    cd.add_series("P50 [ms]", (533, 279, 1943, 2519))
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.35), Inches(1.6), Inches(7.4), Inches(3.6), cd)
    ch = gf.chart
    ch.has_legend = False
    ser = ch.plots[0].series[0]
    ch.plots[0].gap_width = 55
    for i, c in enumerate([GREEN, RGBColor(0x9A, 0xA4, 0xB8), NAVY, RED]):
        pt_ = ser.points[i]
        pt_.format.fill.solid()
        pt_.format.fill.fore_color.rgb = c
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.number_format = '#,##0" ms"'
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(11)
    pl.data_labels.font.bold = True
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = 3000, 0
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE3, 0xE3)
    va.tick_labels.font.size = Pt(10)
    va.tick_labels.font.color.rgb = GRAY
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.reverse_order = True
    caption(s, 0.35, 5.25, 7.4, 16, "chunk 한 번의 지연 P50 (T0 측정, 24 chunk) — 24/24 가 예산 초과, AG3S 가 서버 시간의 77 %", pt=10)
    table(s, 0.35, 5.75, [3.1, 2.1, 2.2],
          [["camera 시차 → 판정 정지(HOLD)", "서버 렌더링", "로컬 GPU 렌더링"],
           ["HOLD chunk (75 개 중)", "42", "0"]], pt=10.5, row_h=0.34, first_col_bold=True,
          col_align=[None, PP_ALIGN.CENTER, PP_ALIGN.CENTER])
    y = text_block(s, 8.1, 1.62, 4.9, [
        section_head("camera 시차 — 해결", pt=13),
        body("서버는 세 camera 를 차례로 렌더링해 촬영 시각이 97–111 ms 벌어졌다", pt=11.5),
        body("시각이 어긋난 관측은 '기하 미인증' → 안전을 위해 HOLD", pt=11.5),
        body("관측 하나에 촬영 시각 하나로 바꿔 시차 0 ms", pt=11.5, after=8),
        section_head("실시간 — 보류 (판단의 정확성 먼저)", pt=13),
        body("지배 항은 AG3S (P50 1943 ms) — 단계별 profiling 이 아직 없다", pt=11.5),
        body("서버 통신은 왕복의 7.6 % 뿐 — 옮겨도 예산을 못 맞춘다", pt=11.5),
        body("control 이 쓰는 거리장은 평균 2.6 초 전 것이다", pt=11.5),
    ], name="timing text")
    st = ("여덟 번째는 시간 문제입니다. 8 step 을 실행하는 533 ms 안에 지각, 거리장, 최적화를 끝내야 하는데, 지금 서버 전체는 "
          "중앙값 2.5 초, 예산의 4.7 배가 걸립니다. 24 개 chunk 전부 초과했고, 그중 77 % 가 AG3S 입니다.\n\n"
          "시간 문제는 두 갈래였습니다. 하나는 camera 시차입니다. 서버가 세 camera 를 차례로 렌더링하다 보니 촬영 시각이 약 "
          "100 ms 벌어졌고, 시각이 어긋난 관측은 기하를 믿을 수 없다고 보고 로봇을 멈췄습니다. 서버에서는 75 chunk 중 42 번 "
          "멈췄는데 로컬 GPU 에서는 0 번이었습니다. 관측 하나에 촬영 시각 하나로 바꿔 시차를 0 으로 만들었습니다.\n\n"
          "다른 하나는 실시간성 자체인데, 지금 단계에서는 판단의 정확성을 먼저 보기로 해서 보류했습니다. 통신은 전체의 8 % 가 "
          "안 되므로 구조를 옮기는 것으로는 해결되지 않고, AG3S 단계별 profiling 부터 해야 합니다.")
    set_notes(s, st)
    SCRIPT.append(("5. Experiment – ⑧ 시간", st))
    ex.append(s)

    # 실험 슬라이드들을 요약 뒤, Future 앞으로 — 현재 순서에서 Future(S[12]) 의 위치를 찾아 그 앞에 끼운다
    ids = list(prs.slides._sldIdLst)

    def index_of(slide):
        return next(i for i, e in enumerate(prs.slides._sldIdLst) if prs.part.related_part(e.rId) is slide.part)

    for sl_ in [sm] + ex:
        move_slide(prs, sl_, index_of(S[12]))

    # Future
    fu = S[12]
    set_title(fu, "6. Future", font=FONT_XB)
    text_block(fu, 0.35, 0.98, 12.6, [{"text": "closed loop 성공률을 VLA 단독 이상으로 — 그다음이 실시간",
                                         "pt": 15, "bold": True, "color": NAVY}], name="future headline")
    items = [
        ("1. Closed loop 성공률", ["관절 한계를 실제 로봇 범위로 · SQP 반복 예산 확보", "N ≥ 10 반복으로 closed loop ≥ shadow 검증"]),
        ("2. Target / Obstacle 분리", ["이웃 과일과 병합된 cluster 를 attention 으로 나누기", "한쪽에서만 본 Target 의 뒷면까지 제외"]),
        ("3. 쥔 물체 처리", ["쥔 사과가 거리장에 남아 손가락과 충돌 → 로봇 모델에 붙여 self-filter 가 지우게"]),
        ("4. 안전 판정 검증", ["거리장의 허용 나이(max_field_age) 결정", "고장을 주입했을 때 정말 멈추는가 (fail-closed)"]),
        ("5. 거리장 구조", ["미세 창 경계 불연속 · self-collision 없음", "사라진 물체의 잔상 제거 (frustum-aware decay)"]),
        ("6. 실시간화", ["AG3S 단계별 profiling 부터", "목표: chunk 예산 533 ms"]),
    ]
    for i, (h, lines) in enumerate(items):
        cx = 0.35 + (i % 3) * 4.25
        cy = 1.65 + (i // 3) * 2.55
        rounded(fu, cx, cy, 4.05, 2.3, fill=LIGHT)
        text_block(fu, cx + 0.18, cy + 0.16, 3.7,
                   [{"text": h, "pt": 16, "bold": True, "color": NAVY, "after": 8}] +
                   [body(t, pt=12.5, after=6) for t in lines], name=f"future/{h}", h=2.05)
    sf = ("앞으로의 계획입니다. 가장 먼저 closed loop 성공률을 VLA 단독 이상으로 올리는 것입니다. 관절 한계를 실제 로봇 범위로 "
          "맞추고 최적화 반복 시간을 확보한 뒤, 10 번 이상 반복해서 closed loop 이 shadow 보다 나쁘지 않음을 보이겠습니다.\n\n"
          "지각 쪽에서는 사과가 이웃 과일과 한 덩어리로 묶이는 경우를 attention 으로 나누는 것, 한쪽에서만 본 Target 의 뒷면까지 "
          "제외하는 것, 그리고 쥔 사과가 거리장에 남아 손가락과 충돌하는 문제를 로봇 모델에 붙여 해결하는 것이 남았습니다.\n\n"
          "안전 판정 쪽에서는 거리장을 몇 초까지 믿을지 정하고, 고장을 넣었을 때 정말 멈추는지 시험해야 합니다. 거리장 구조의 "
          "남은 문제들과, 마지막으로 실시간화가 그다음 순서입니다. 감사합니다.")
    set_notes(fu, sf)
    SCRIPT.append(("6. Future", sf))
    SCRIPT.append(("Thank you / Q&A", "(질문을 받는다)"))

    prs.save(str(OUT))

    # 대본 md — 발표 순서대로
    order = [prs.part.related_part(e.rId) for e in prs.slides._sldIdLst]
    lines = ["# 곽민기 0930 세미나 — 발표 대본", "",
             f"- deck: `{OUT.name}` (같은 대본이 슬라이드마다 speaker notes 에 들어 있다)",
             "- 숫자 출처: `AG3S/ag3s/docs/AG3S_T0T6_LOG.md` · `handoff/AG3S_GRASP_FIX_PLAN.md` · "
             "`handoff/T29.impl.md`(자가 점검) · `handoff/T30.verify.json` · 각 그림 원본",
             "- 그림: `AG3S/ag3s/docs/figures/ppt-0930/` (`ppt-0930-figures.json` 에 원본 · crop 영역)", ""]
    for n, part in enumerate(order, 1):
        slide = next(s_ for s_ in prs.slides if s_.part is part)
        title = (slide.shapes.title.text_frame.text if slide.shapes.title is not None
                 else ("표지" if n == 1 else "Thank you / Q&A"))
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        lines += [f"## {n}. {title}", "", notes or "_(대본 없음)_", ""]
    SCRIPT_MD.write_text("\n".join(lines))
    print("saved", OUT)
    print("\n".join(WARN) or "no overflow warnings")


if __name__ == "__main__":
    main()
