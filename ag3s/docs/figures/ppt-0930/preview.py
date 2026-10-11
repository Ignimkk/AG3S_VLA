"""LibreOffice 없이 슬라이드 배치를 보는 근사 렌더러 (겹침 · 넘침 확인용).

글꼴은 Noto Sans CJK 로 대신 그린다(NanumSquare Neo 폭과 비슷하지만 같지 않다).
표 · 그림 · 도형 · 글상자와 layout/master 의 그림을 위치대로 그린다. chart 는 회색 상자로 대신한다.
"""
import sys
import pathlib

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

DPI = 110
REG = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"


def px(v):
    return int(Emu(v).inches * DPI)


def font(pt, bold):
    return ImageFont.truetype(BOLD if bold else REG, max(6, int(pt * DPI / 72)))


def draw_text_frame(d, tf, x, y, w, h, default_pt=18):
    cy = y + 4
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            cy += int(default_pt * DPI / 72 * 1.2)
            continue
        pt = (runs[0].font.size.pt if runs[0].font.size else default_pt)
        bold = bool(runs[0].font.bold)
        pPr = p._p.pPr
        indent = 0
        bullet = ""
        if pPr is not None:
            marL = pPr.get("marL")
            if marL:
                indent = px(int(marL))
            if pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None:
                bullet = "•"
            if pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum") is not None:
                bullet = "#"
        f = font(pt, bold)
        lh = int(pt * DPI / 72 * 1.25)
        segs = []
        for r in runs:
            col = (40, 40, 40)
            try:
                if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
                    c = r.font.color.rgb
                    col = (c[0], c[1], c[2])
            except Exception:
                pass
            rb = bool(r.font.bold) if r.font.bold is not None else bold
            rpt = r.font.size.pt if r.font.size else pt
            segs.append((r.text, font(rpt, rb), col))
        if bullet:
            d.text((x + indent - int(12 * DPI / 72), cy), bullet, font=f, fill=(40, 40, 40))
        cx = x + indent
        limit = x + w - 4
        for txt, ff, col in segs:
            for ch in txt:
                cw = ff.getlength(ch)
                if cx + cw > limit:
                    cy += lh
                    cx = x + indent
                d.text((cx, cy), ch, font=ff, fill=col)
                cx += cw
        cy += lh + int((p.space_after.pt if p.space_after else 0) * DPI / 72)
    if cy > y + h + 6:
        d.rectangle([x, y, x + w, y + h], outline=(255, 0, 0), width=2)
        return cy - (y + h)
    return 0


def render(path, outdir, only=None):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    outs = []
    for n, slide in enumerate(prs.slides, 1):
        if only and n not in only:
            continue
        im = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(im)
        over = []
        layers = [slide.slide_layout.slide_master.shapes, slide.slide_layout.shapes, slide.shapes]
        for li, shapes in enumerate(layers):
            for sh in shapes:
                if li < 2 and sh.is_placeholder:
                    continue
                if sh.left is None:
                    continue
                x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        pic = Image.open(sh.image._blob if False else __import__("io").BytesIO(sh.image.blob)).convert("RGB")
                        im.paste(pic.resize((max(1, w), max(1, h))), (x, y))
                    except Exception:
                        d.rectangle([x, y, x + w, y + h], outline=(0, 0, 255))
                elif sh.has_chart if hasattr(sh, "has_chart") else False:
                    d.rectangle([x, y, x + w, y + h], fill=(235, 238, 245), outline=(150, 150, 150))
                    d.text((x + 10, y + 10), "[chart]", font=font(12, True), fill=(90, 90, 90))
                elif sh.has_table if hasattr(sh, "has_table") else False:
                    tbl = sh.table
                    cy = y
                    for r_i, row in enumerate(tbl.rows):
                        rh = px(row.height)
                        cx = x
                        for c_i, col in enumerate(tbl.columns):
                            cw = px(col.width)
                            cell = tbl.cell(r_i, c_i)
                            fill = (255, 255, 255)
                            try:
                                if cell.fill.type is not None:
                                    c = cell.fill.fore_color.rgb
                                    fill = (c[0], c[1], c[2])
                            except Exception:
                                pass
                            d.rectangle([cx, cy, cx + cw, cy + rh], fill=fill, outline=(200, 200, 200))
                            o = draw_text_frame(d, cell.text_frame, cx + 4, cy, cw - 8, rh, 10)
                            if o:
                                over.append(f"table cell r{r_i}c{c_i} +{o}px")
                            cx += cw
                        cy += rh
                else:
                    try:
                        if sh.fill.type == 1:
                            c = sh.fill.fore_color.rgb
                            d.rounded_rectangle([x, y, x + w, y + h], radius=8, fill=(c[0], c[1], c[2]))
                    except Exception:
                        pass
                    if sh.shape_type == MSO_SHAPE_TYPE.LINE:
                        d.line([x, y, x + w, y + h], fill=(0, 0, 0), width=1)
                    if sh.has_text_frame and sh.text_frame.text.strip():
                        o = draw_text_frame(d, sh.text_frame, x, y, w, h,
                                            24 if sh.is_placeholder else 18)
                        if o and not sh.is_placeholder:
                            over.append(f"{sh.name!r} '{sh.text_frame.text[:30]}' +{o}px")
        p = pathlib.Path(outdir) / f"s{n:02d}.png"
        im.save(p)
        outs.append((n, p, over))
    return outs


if __name__ == "__main__":
    only = set(int(a) for a in sys.argv[3:]) if len(sys.argv) > 3 else None
    for n, p, over in render(sys.argv[1], sys.argv[2], only):
        print(n, p, over if over else "")
